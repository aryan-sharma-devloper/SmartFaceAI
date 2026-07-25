from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import (
    Font,
    PatternFill,
    Alignment,
    Border,
    Side,
)
from openpyxl.utils import get_column_letter


def generate_attendance_excel(records):

    wb = Workbook()
    ws = wb.active
    ws.title = "Attendance Report"

    # ------------------------------------------------
    # Title
    # ------------------------------------------------

    ws.merge_cells("A1:H1")

    title = ws["A1"]
    title.value = "SMART FACE AI - ATTENDANCE REPORT"

    title.font = Font(
        size=18,
        bold=True,
        color="FFFFFF",
    )

    title.fill = PatternFill(
        fill_type="solid",
        fgColor="1E40AF",
    )

    title.alignment = Alignment(
        horizontal="center",
        vertical="center",
    )

    ws.row_dimensions[1].height = 30

    # ------------------------------------------------
    # Headers
    # ------------------------------------------------

    headers = [
        "Employee ID",
        "Employee Name",
        "Department",
        "Date",
        "Check In",
        "Check Out",
        "Working Hours",
        "Status",
    ]

    header_fill = PatternFill(
        fill_type="solid",
        fgColor="2563EB",
    )

    header_font = Font(
        color="FFFFFF",
        bold=True,
    )

    thin = Side(
        style="thin",
        color="D1D5DB",
    )

    border = Border(
        left=thin,
        right=thin,
        top=thin,
        bottom=thin,
    )

    header_row = 3

    for col, header in enumerate(headers, start=1):

        cell = ws.cell(
            row=header_row,
            column=col,
        )

        cell.value = header
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )
        cell.border = border

    # ------------------------------------------------
    # Data
    # ------------------------------------------------

    row = header_row + 1

    for record in records:

        values = [
            record["employee_id"],
            record["full_name"],
            record["department"],
            str(record["date"]),
            str(record["check_in"]),
            str(record["check_out"]),
            str(record["working_hours"]),
            record["status"],
        ]

        for col, value in enumerate(values, start=1):

            cell = ws.cell(
                row=row,
                column=col,
            )

            cell.value = value
            cell.border = border
            cell.alignment = Alignment(
                horizontal="center",
            )

            # Alternate row color
            if row % 2 == 0:
                cell.fill = PatternFill(
                    fill_type="solid",
                    fgColor="F8FAFC",
                )

            # Status colors
            if col == 8:

                status = str(value).lower()

                if status == "present":
                    cell.fill = PatternFill(
                        fill_type="solid",
                        fgColor="DCFCE7",
                    )

                elif status == "late":
                    cell.fill = PatternFill(
                        fill_type="solid",
                        fgColor="FEF3C7",
                    )

                elif status == "absent":
                    cell.fill = PatternFill(
                        fill_type="solid",
                        fgColor="FECACA",
                    )

        row += 1

    # ------------------------------------------------
    # Auto Width
    # ------------------------------------------------

    for column_cells in ws.columns:

        max_length = 0
        column = column_cells[0].column

        for cell in column_cells:

            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass

        adjusted_width = max_length + 5

        ws.column_dimensions[
            get_column_letter(column)
        ].width = adjusted_width

    output = BytesIO()

    wb.save(output)

    output.seek(0)

    return output